#!/usr/bin/env python3
"""生成原布局一页洞察 PPT：无顶/底黑底、无绿色填充，仅边框，强调色用暗红色。"""
from __future__ import annotations

from pathlib import Path

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import PP_ALIGN
from pptx.oxml import parse_xml
from pptx.oxml.ns import qn
from pptx.util import Inches, Pt

ROOT = Path(__file__).resolve().parent
OUT = ROOT / "docs"
ART = Path("/opt/cursor/artifacts/yunqi-docs")

# 原布局浅底 + 暗红色强调（无黑底、无绿底）
BG = RGBColor(0xF5, 0xF7, 0xF9)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)
DARK = RGBColor(0x1A, 0x1A, 0x1A)
MID = RGBColor(0x4A, 0x4A, 0x4A)
MUTED = RGBColor(0x6B, 0x6B, 0x6B)
BORDER = RGBColor(0xD0, 0xD8, 0xDE)
ACCENT = RGBColor(0x8B, 0x1A, 0x2B)  # 暗红色
FONT = "Microsoft YaHei"

W = Inches(13.333)
H = Inches(7.5)


def set_run(run, size, bold=False, color=DARK):
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.color.rgb = color
    run.font.name = FONT
    rPr = run._r.get_or_add_rPr()
    ea = rPr.find(qn("a:ea"))
    if ea is None:
        ea = parse_xml(
            f'<a:ea xmlns:a="http://schemas.openxmlformats.org/drawingml/2006/main" typeface="{FONT}"/>'
        )
        rPr.append(ea)
    else:
        ea.set("typeface", FONT)


def rect(slide, x, y, w, h, fill=None, line=None, line_w=1.0):
    sh = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, x, y, w, h)
    if fill is None:
        sh.fill.background()
    else:
        sh.fill.solid()
        sh.fill.fore_color.rgb = fill
    if line is None:
        sh.line.fill.background()
    else:
        sh.line.color.rgb = line
        sh.line.width = Pt(line_w)
    return sh


def textbox(slide, x, y, w, h, text, size=14, bold=False, color=DARK, align=PP_ALIGN.LEFT):
    box = slide.shapes.add_textbox(x, y, w, h)
    tf = box.text_frame
    tf.word_wrap = True
    p = tf.paragraphs[0]
    p.alignment = align
    run = p.add_run()
    run.text = text
    set_run(run, size=size, bold=bold, color=color)
    return box


def multilines(slide, x, y, w, h, lines, size=13, color=MID, spacing=4, bold_first=False):
    box = slide.shapes.add_textbox(x, y, w, h)
    tf = box.text_frame
    tf.word_wrap = True
    for i, line in enumerate(lines):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = PP_ALIGN.LEFT
        p.space_after = Pt(spacing)
        run = p.add_run()
        run.text = line
        set_run(run, size=size, bold=(bold_first and i == 0), color=DARK if (bold_first and i == 0) else color)
    return box


def build_one_pager(
    filename: str,
    eyebrow: str,
    insight: str,
    judgment: str,
    cards: list[tuple[str, str, list[str]]],
    metrics: list[tuple[str, str, str]],
    action: str,
    source: str,
    metrics_title: str = "关键指标对比与提升",
):
    prs = Presentation()
    prs.slide_width = W
    prs.slide_height = H
    slide = prs.slides.add_slide(prs.slide_layouts[6])

    # 页面浅底
    rect(slide, 0, 0, W, H, fill=BG)

    # 左侧仅暗红色细边（原绿条改为边框线）
    rect(slide, 0, 0, Pt(4), H, fill=ACCENT)

    # 顶部：无黑底，仅暗红色底边框区域
    header_h = Inches(1.15)
    rect(slide, 0, 0, W, header_h, fill=WHITE, line=None)
    rect(slide, 0, header_h - Pt(2), W, Pt(2), fill=ACCENT)
    textbox(slide, Inches(0.4), Inches(0.16), Inches(12.5), Inches(0.3), eyebrow, size=13, color=MUTED)
    textbox(slide, Inches(0.4), Inches(0.46), Inches(12.5), Inches(0.55), insight, size=22, bold=True, color=DARK)

    # 一句话判断：白底 + 暗红色边框（原浅绿底/绿边）
    jy = Inches(1.32)
    rect(slide, Inches(0.35), jy, Inches(12.6), Inches(0.55), fill=WHITE, line=ACCENT, line_w=1.5)
    textbox(slide, Inches(0.5), jy + Inches(0.1), Inches(12.3), Inches(0.4), judgment, size=14, bold=True, color=DARK)

    # 三栏卡片：白底灰边，标题区无绿底，仅顶边暗红色细线
    cy = Inches(2.05)
    ch = Inches(2.85)
    cw = Inches(4.05)
    gap = Inches(0.22)
    for i, (title, lead, bullets) in enumerate(cards):
        x = Inches(0.35) + i * (cw + gap)
        rect(slide, x, cy, cw, ch, fill=WHITE, line=BORDER, line_w=1)
        rect(slide, x, cy, cw, Pt(3), fill=ACCENT)  # 顶边强调线，非大块填充
        textbox(slide, x + Inches(0.15), cy + Inches(0.15), cw - Inches(0.3), Inches(0.35), title, size=16, bold=True, color=ACCENT)
        lines = [lead] + bullets
        multilines(
            slide,
            x + Inches(0.18),
            cy + Inches(0.55),
            cw - Inches(0.36),
            ch - Inches(0.7),
            lines,
            size=13,
            color=MID,
            spacing=5,
            bold_first=True,
        )

    # 指标区标题
    my = Inches(5.05)
    textbox(slide, Inches(0.35), my, Inches(8), Inches(0.28), metrics_title, size=14, bold=True, color=DARK)

    # 指标卡：白底边框 + 左侧暗红色细条（原橙色）
    mw = Inches(2.02)
    mh = Inches(0.95)
    mg = Inches(0.12)
    for i, (label, baseline, value) in enumerate(metrics):
        x = Inches(0.35) + i * (mw + mg)
        y = my + Inches(0.32)
        rect(slide, x, y, mw, mh, fill=WHITE, line=BORDER, line_w=1)
        rect(slide, x, y, Pt(4), mh, fill=ACCENT)
        textbox(slide, x + Inches(0.16), y + Inches(0.06), mw - Inches(0.22), Inches(0.22), label, size=11, color=MUTED)
        textbox(slide, x + Inches(0.16), y + Inches(0.28), mw - Inches(0.22), Inches(0.28), baseline, size=11, color=MID)
        textbox(slide, x + Inches(0.16), y + Inches(0.55), mw - Inches(0.22), Inches(0.32), value, size=18, bold=True, color=ACCENT)

    # 底部：无黑底，仅顶边暗红色细线 + 白底
    fy = Inches(6.55)
    rect(slide, 0, fy, W, H - fy, fill=WHITE, line=None)
    rect(slide, 0, fy, W, Pt(2), fill=ACCENT)
    textbox(slide, Inches(0.4), fy + Inches(0.1), Inches(12.5), Inches(0.28), action, size=12, bold=True, color=DARK)
    textbox(slide, Inches(0.4), fy + Inches(0.42), Inches(12.5), Inches(0.28), source, size=10, color=MUTED)

    OUT.mkdir(parents=True, exist_ok=True)
    ART.mkdir(parents=True, exist_ok=True)
    out = OUT / filename
    prs.save(out)
    prs.save(ART / filename)
    return out


def main():
    outs = []
    outs.append(
        build_one_pager(
            "洞察一页-04-KVCache推理基础设施.pptx",
            "04 · 王正恒 / 徐国强 · KVCache：从显存优化到软硬结合的推理基础设施",
            "洞察：KVCache 已从框架技巧，升级为可调度的推理基础设施",
            "一句话判断：把 KV 当资产——控制面管放置/生命周期，数据面做分层池化；无仿真就没有可规模化的 SLO/TCO。",
            [
                (
                    "01  三跃迁：资产化",
                    "从显存开销 → 可复用推理资产",
                    [
                        "• 前缀命中：跳过 Prefill，降 TTFT",
                        "• Agentic 多轮：KV 跨轮持久",
                        "• 跨实例池化：容量与计算解耦",
                        "• KV Aware Routing 感知放置",
                    ],
                ),
                (
                    "02  分层栈：破存储墙",
                    "G1→G4 是标配，不是可选项",
                    [
                        "• G1 HBM / G1.5 UMX·AMES",
                        "• G2 DRAM / G2.5 CXL·RDMA 池",
                        "• G3 SSD / G3.5 EBoF / G4 对象",
                        "• UMX：UALink+CXL Load/Store",
                    ],
                ),
                (
                    "03  决策：仿真驱动",
                    "经验调参 → HiSim Pareto/SLO",
                    [
                        "• KVCM 控制面 + MemPool 数据面",
                        "• 带宽/容量/寿命三约束取最大",
                        "• Token 级 TCO 进入平台指标",
                        "• 与 CIPU/引擎纵向对齐 L1–L4",
                    ],
                ),
            ],
            [
                ("offload 吞吐", "未池化 → MemPool", "+50%+"),
                ("TTFT", "未 offload → 池化", "-30%"),
                ("传输效率", "多网卡数据面", "93%+"),
                ("Token TCO", "DRAM 为主 → 分层", "降幅>80%*"),
                ("HiSim 误差", "经验配置 → 仿真", "<3%*"),
                ("分层形态", "单机 HBM", "G1–G4 栈"),
            ],
            "可落地动作：把命中率 / TTFT / $/Token 写入平台 SLO；用 workload 仿真选型 CXL/SSD/EBoF；打通 KVCM 与引擎路由。",
            "*为演讲幻灯宣称口径  ·  来源：云栖大会 agendaId=164 · Talk 04 王正恒 / 徐国强",
        )
    )
    outs.append(
        build_one_pager(
            "洞察一页-07-AI推理全栈优化实战.pptx",
            "07 · 资彦义 / 李陈浩文 · Agentic 时代 AI 推理软硬件结合全栈优化实战",
            "洞察：全栈优化三件套 —— Agentic KV · Mega Kernel · Kernel Agent",
            "一句话判断：停得久≠价值低；微秒尺度看 kernel 之间；跨芯竞争力取决于算子供给能否 Agent 化。",
            [
                (
                    "01  Agentic KV",
                    "LRU → 会话价值调度",
                    [
                        "• 暂停会话常被 LRU 误删",
                        "• 生命周期 + 重建成本/命中",
                        "• 绿保留 / 黄准备 / 红腾挪",
                        "• 避免强制 Prefill 吞噬 SLA",
                    ],
                ),
                (
                    "02  Mega Kernel",
                    "算子接力 → 设备内流水",
                    [
                        "• Host-driven → Device-resident",
                        "• 少启动 · 少搬运 · 少等待",
                        "• 同步粒度：完成 → 数据就绪",
                        "• Fusion→SM Task→Runtime→Kernel",
                    ],
                ),
                (
                    "03  Kernel Agent",
                    "专家手艺 → 可迁移资产",
                    [
                        "• Master–Sub 并行探索回写",
                        "• H100 案例约 4.88×",
                        "• PPU / AMD 跨芯迁移加速",
                        "• Dev Loop：Profiler→回归门禁",
                    ],
                ),
            ],
            [
                ("H100 BF16", "9.52 → 1.95 ms", "4.88×"),
                ("PPU MoE gate", "相对 SGLang", "1.56–2.13×"),
                ("PPU 小算子链", "RMSNorm/SiLU", "1.23–1.29×"),
                ("AMD Decode", "DSA 路径", "约 1.6×"),
                ("AMD GDN", "递推核", "约 1.18×"),
                ("缓存策略", "LRU → 价值排序", "少 Prefill"),
            ],
            "可落地动作：工程化 Dev Loop（Trace+回归）；KV 调度接入会话状态并与池化联动；多芯导入优先建 Kernel Agent+评测集。",
            "来源：云栖大会 agendaId=164 · Talk 07 资彦义 / 李陈浩文",
        )
    )
    outs.append(
        build_one_pager(
            "洞察一页-08-未来大模型推理架构.pptx",
            "08 · 李文韬 · 未来大模型推理的软硬件结合系统架构探索",
            "洞察：先拆分，再堆料；先仿真，再到柜",
            "一句话判断：P/D + AFD 让 Attention/KV 与 FFN/Expert 各得其所；无高保真仿真，A/F 配比无法在到货前收敛。",
            [
                (
                    "01  错配：为何要拆",
                    "模型与硬件不同步增长",
                    [
                        "• 上下文：128K → ~1M 级",
                        "• MoE：总参可 2.8T / 激活~95B",
                        "• Prefill / Attn / FFN 画像各异",
                        "• 单卡堆料无法对齐三阶段",
                    ],
                ),
                (
                    "02  拆分：P/D + AFD",
                    "阶段级 + 层内算子级可组合",
                    [
                        "• P/D：Prefill 与 Decode 解耦",
                        "• AFD：Attention 池 ↔ Expert 池",
                        "• A 卡：大 HBM + 高 KV 带宽",
                        "• F 卡：LPU 矩阵 + 高权重带宽",
                    ],
                ),
                (
                    "03  前移：Insight/Hisim",
                    "无实物仿真成为导入标配",
                    [
                        "• 保留 SGLang 调度/组批逻辑",
                        "• 虚拟时钟替代真实 Forward",
                        "• 整体准确率 96%+",
                        "• 开源合入 → AFD PR → POC",
                    ],
                ),
            ],
            [
                ("上下文", "128K → ~1M", "长上下文压"),
                ("MoE 形态", "2.8T / 激活~95B", "稀激活"),
                ("LPU 极端点", "高带宽低容量", "FFN 侧位"),
                ("Hisim", "无实物选型", "96%+"),
                ("架构", "同构堆叠", "P/D+AFD"),
                ("落地路径", "仿真→PR→POC", "先仿真"),
            ],
            "可落地动作：并行推进拆分架构与仿真器；跟踪 HiSim/AFD 开源合入；用业务模型先跑配比，再定超节点与 A/F 采购。",
            "来源：云栖大会 agendaId=164 · Talk 08 李文韬  ·  与 Talk04/06/07 形成「分层 + 超节点 + 全栈 + 拆分」闭环",
        )
    )
    outs.append(
        build_one_pager(
            "洞察一页-Agentic-AI-Infra.pptx",
            "2026 云栖大会 · Agentic AI 算力与存储论坛",
            "洞察：瓶颈已从「堆卡」转向「拆分 · 分层 · 卸载 · 平台化」",
            "一句话判断：Agentic AI 让系统胜负手变成 安全隔离 + KV/记忆路径 + 调度成熟度 + 架构拆分；堆算力收益递减。",
            [
                (
                    "01  瓶颈迁移",
                    "从 GPU 峰值算力 → 数据访问 / 记忆 / 安全 / 调度",
                    [
                        "• 算力廉价、访存昂贵（CIPU）",
                        "• 停得久 ≠ 价值低（Agentic KV）",
                        "• 微秒尺度看 kernel 之间（Mega Kernel）",
                        "• 运维先建评测床，再谈智能体",
                    ],
                ),
                (
                    "02  分层与拆分",
                    "垂直分层 + 水平拆分，替代单卡堆料",
                    [
                        "• KV：G1 HBM → UMX/CXL → SSD/EBoF",
                        "• 存储：按训练/推理/Agentic 分温度",
                        "• 架构：P/D + AFD（A卡 / LPU F卡）",
                        "• 控制面与数据面必须拆开演进",
                    ],
                ),
                (
                    "03  平台化工程",
                    "把专家手艺变成可复制系统",
                    [
                        "• 震旦：交付 / 仿真 / 诊断闭环",
                        "• HiSim/Insight：选型前移（96%+）",
                        "• Kernel Agent：跨芯算子自动供给",
                        "• Infini-AIOps：工单时长约 -91%",
                    ],
                ),
            ],
            [
                ("运维时长", "5.5h → 0.47h", "-91%"),
                ("KV 吞吐 / TTFT", "+50%+ / -30%", "MemPool"),
                ("AL144 vs 128", "卡数·功率", "2× / 2.15×"),
                ("Kernel Agent", "9.52→1.95ms", "4.88×"),
                ("CPU:GPU", "1:8 → 1:1~2", "核数~4×"),
                ("Hisim", "仿真准确率", "96%+"),
            ],
            "近半年可落地：KV 分层与价值调度 · 运维评测床 · Profiler Dev Loop · 安全八项清单",
            "中远期：CXL/SCM 与超节点功率规划 · Kernel Agent 多芯供给 · P/D+AFD 拆分（先仿真再到柜）  |  来源：云栖大会 agendaId=164 八场技术分享综合",
            metrics_title="关键指标对比（可直接做对比页）",
        )
    )
    for p in outs:
        print("wrote", p)


if __name__ == "__main__":
    main()
