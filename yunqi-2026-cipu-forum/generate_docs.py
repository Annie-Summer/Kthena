#!/usr/bin/env python3
"""Generate Word summaries for Yunqi 2026 CIPU forum talks."""
from __future__ import annotations

import json
from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor

ROOT = Path("/workspace/yunqi-2026-cipu-forum")
NOTES = ROOT / "notes"
DOCS = ROOT / "docs"
SHOTS = ROOT / "screenshots"

INSIGHT_LABELS = [
    ("trend", "趋势判断"),
    ("tech_judgment", "关键技术判断"),
    ("infra_implication", "对 Infra / 平台的启示"),
    ("ppt_hooks", "洞察PPT页建议（可直接拆页）"),
]


def set_run_font(run, name="微软雅黑", size=11, bold=False, color=None):
    run.font.name = name
    run._element.rPr.rFonts.set(qn("w:eastAsia"), name)
    run.font.size = Pt(size)
    run.bold = bold
    if color:
        run.font.color.rgb = RGBColor(*color)


def add_heading(doc, text, level=1):
    p = doc.add_heading(text, level=level)
    for run in p.runs:
        set_run_font(run, size=16 if level == 1 else 13, bold=True)
    return p


def add_para(doc, text, *, bold=False, size=11, space_after=6):
    p = doc.add_paragraph()
    run = p.add_run(text)
    set_run_font(run, size=size, bold=bold)
    p.paragraph_format.space_after = Pt(space_after)
    p.paragraph_format.line_spacing = 1.25
    return p


def add_bullets(doc, items):
    for item in items:
        p = doc.add_paragraph(style="List Bullet")
        run = p.add_run(item)
        set_run_font(run, size=11)
        p.paragraph_format.space_after = Pt(3)


def add_image(doc, path: Path, caption: str | None = None, width=5.8):
    if not path.exists():
        add_para(doc, f"[缺少截图: {path.name}]", bold=True)
        return
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run()
    run.add_picture(str(path), width=Inches(width))
    if caption:
        cp = doc.add_paragraph()
        cp.alignment = WD_ALIGN_PARAGRAPH.CENTER
        r = cp.add_run(caption)
        set_run_font(r, size=9, color=(90, 90, 90))


def add_evidence_images(doc, shot_dir: str, images: list[dict]):
    for img in images:
        img_dir = img.get("shot_dir") or shot_dir
        path = SHOTS / img_dir / img["file"]
        add_image(doc, path, caption=img.get("caption"))
        if img.get("note"):
            add_para(doc, f"佐证说明：{img['note']}", size=10, space_after=10)


def add_metrics_table(doc, rows: list[dict]):
    if not rows:
        return
    add_para(
        doc,
        "下表汇总本场（或跨场）可核对的关键指标：基线 → 改进后 → 提升幅度。后续洞察PPT可直接做成对比页。",
        size=10,
    )
    table = doc.add_table(rows=1 + len(rows), cols=5)
    table.style = "Table Grid"
    headers = ["指标", "基线/对照", "改进后/方案", "提升/变化", "出处"]
    for i, h in enumerate(headers):
        cell = table.rows[0].cells[i]
        cell.text = h
        for p in cell.paragraphs:
            for run in p.runs:
                set_run_font(run, size=10, bold=True)
        tc = cell._tc
        tcPr = tc.get_or_add_tcPr()
        shd = OxmlElement("w:shd")
        shd.set(qn("w:fill"), "D9E2F3")
        shd.set(qn("w:val"), "clear")
        tcPr.append(shd)

    for r_idx, row in enumerate(rows, start=1):
        vals = [
            row.get("name", ""),
            row.get("baseline", ""),
            row.get("improved", ""),
            row.get("delta", ""),
            row.get("source", ""),
        ]
        for c_idx, val in enumerate(vals):
            cell = table.rows[r_idx].cells[c_idx]
            cell.text = val
            for p in cell.paragraphs:
                for run in p.runs:
                    set_run_font(run, size=9, bold=(c_idx == 3))
    doc.add_paragraph()


def build_doc(meta: dict) -> Path:
    doc = Document()
    section = doc.sections[0]
    section.top_margin = Inches(0.8)
    section.bottom_margin = Inches(0.8)
    section.left_margin = Inches(0.9)
    section.right_margin = Inches(0.9)

    add_heading(doc, meta["title"], 1)
    add_para(doc, f"演讲嘉宾：{meta['speakers']}", bold=True)
    add_para(doc, f"大会议程时间：{meta['agenda_time']}")
    add_para(doc, f"视频时间轴：{meta['video_time']}")
    add_para(doc, "来源：2026云栖大会 · Agentic AI 时代的算力与存储服务器革新")
    add_para(doc, "回放链接：https://yunqi.aliyun.com/2026/session?agendaId=164")
    add_para(
        doc,
        "文档定位：偏技术细节与未来启示，可作为后续洞察PPT的素材底稿。",
        size=10,
        bold=True,
    )

    add_heading(doc, "一、议题速览", 2)
    add_para(doc, meta["summary"])

    if meta.get("tech_highlights"):
        add_heading(doc, "二、关键技术细节速览", 2)
        add_bullets(doc, meta["tech_highlights"])

    if meta.get("metrics"):
        add_heading(doc, "三、关键技术指标对比与提升", 2)
        add_metrics_table(doc, meta["metrics"])

    section_title = "四、技术展开与PPT佐证" if meta.get("metrics") else "三、技术展开与PPT佐证"
    add_heading(doc, section_title, 2)
    add_para(
        doc,
        "以下按技术主题展开；每节要点之后立即附上对应 PPT 截图。",
        size=10,
    )
    for block in meta.get("blocks", []):
        add_para(doc, block["heading"], bold=True, size=12, space_after=4)
        add_bullets(doc, block.get("bullets", []))
        add_evidence_images(doc, meta["shot_dir"], block.get("images", []))

    insights = meta.get("insights") or {}
    if insights:
        insight_title = "五、面向未来的洞察（洞察PPT素材）" if meta.get("metrics") else "四、面向未来的洞察（洞察PPT素材）"
        add_heading(doc, insight_title, 2)
        add_para(
            doc,
            "本节可直接拆成洞察PPT的判断页/对比页/动作页/金句页；指标对比见上文表格。",
            size=10,
        )
        for key, label in INSIGHT_LABELS:
            items = insights.get(key) or []
            if not items:
                continue
            add_para(doc, label, bold=True, size=12, space_after=4)
            add_bullets(doc, items)

    DOCS.mkdir(parents=True, exist_ok=True)
    out = DOCS / meta["filename"]
    doc.save(out)
    return out


def main():
    notes_path = NOTES / "talk_notes.json"
    data = json.loads(notes_path.read_text(encoding="utf-8"))
    outputs = []
    for meta in data:
        path = build_doc(meta)
        outputs.append(str(path))
        print("wrote", path)
    print(json.dumps(outputs, ensure_ascii=False, indent=2))


if __name__ == "__main__":
    main()
