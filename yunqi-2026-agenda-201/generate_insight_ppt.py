#!/usr/bin/env python3
"""为 PAI / Agent Infra 论坛各子议题生成与 CIPU 同风格的一页洞察 PPT（洞察页+截图佐证页）。"""
from __future__ import annotations

import argparse
import json
import re
from pathlib import Path

from PIL import Image
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import PP_ALIGN
from pptx.oxml import parse_xml
from pptx.oxml.ns import qn
from pptx.util import Inches, Pt

BG = RGBColor(0xF5, 0xF7, 0xF9)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)
DARK = RGBColor(0x1A, 0x1A, 0x1A)
MID = RGBColor(0x4A, 0x4A, 0x4A)
MUTED = RGBColor(0x6B, 0x6B, 0x6B)
BORDER = RGBColor(0xD0, 0xD8, 0xDE)
ACCENT = RGBColor(0x8B, 0x1A, 0x2B)
FONT = "Microsoft YaHei"
W = Inches(13.333)
H = Inches(7.5)


def set_run(run, size, bold=False, color=DARK):
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.color.rgb = color
    run.font.name = FONT
    rPr = run._r.get_or_add_rPr()
    ea = rPr.find(qn("a:ea"))
    if ea is None:
        ea = parse_xml(
            f'<a:ea xmlns:a="http://schemas.openxmlformats.org/drawingml/2006/main" typeface="{FONT}"/>'
        )
        rPr.append(ea)
    else:
        ea.set("typeface", FONT)


def rect(slide, x, y, w, h, fill=None, line=None, line_w=1.0):
    sh = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, x, y, w, h)
    if fill is None:
        sh.fill.background()
    else:
        sh.fill.solid()
        sh.fill.fore_color.rgb = fill
    if line is None:
        sh.line.fill.background()
    else:
        sh.line.color.rgb = line
        sh.line.width = Pt(line_w)
    return sh


def textbox(slide, x, y, w, h, text, size=14, bold=False, color=DARK, align=PP_ALIGN.LEFT):
    box = slide.shapes.add_textbox(x, y, w, h)
    tf = box.text_frame
    tf.word_wrap = True
    p = tf.paragraphs[0]
    p.alignment = align
    run = p.add_run()
    run.text = text
    set_run(run, size=size, bold=bold, color=color)
    return box


def multilines(slide, x, y, w, h, lines, size=13, color=MID, spacing=4, bold_first=False):
    box = slide.shapes.add_textbox(x, y, w, h)
    tf = box.text_frame
    tf.word_wrap = True
    for i, line in enumerate(lines):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = PP_ALIGN.LEFT
        p.space_after = Pt(spacing)
        run = p.add_run()
        run.text = line
        set_run(run, size=size, bold=(bold_first and i == 0), color=DARK if (bold_first and i == 0) else color)
    return box


def add_insight_slide(prs, eyebrow, insight, judgment, cards, metrics, action, source, metrics_title="关键指标对比与提升"):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    rect(slide, 0, 0, W, H, fill=BG)
    rect(slide, 0, 0, W, Inches(1.15), fill=WHITE, line=None)
    textbox(slide, Inches(0.4), Inches(0.16), Inches(12.5), Inches(0.3), eyebrow, size=13, color=MUTED)
    textbox(slide, Inches(0.4), Inches(0.46), Inches(12.5), Inches(0.55), insight, size=20, bold=True, color=DARK)

    jy = Inches(1.32)
    rect(slide, Inches(0.35), jy, Inches(12.6), Inches(0.55), fill=WHITE, line=BORDER, line_w=1)
    textbox(slide, Inches(0.5), jy + Inches(0.1), Inches(12.3), Inches(0.4), judgment, size=13, bold=True, color=DARK)

    cy = Inches(2.05)
    ch = Inches(2.85)
    cw = Inches(4.05)
    gap = Inches(0.22)
    for i, (title, lead, bullets) in enumerate(cards[:3]):
        x = Inches(0.35) + i * (cw + gap)
        rect(slide, x, cy, cw, ch, fill=WHITE, line=BORDER, line_w=1)
        textbox(slide, x + Inches(0.15), cy + Inches(0.15), cw - Inches(0.3), Inches(0.35), title, size=15, bold=True, color=ACCENT)
        multilines(
            slide,
            x + Inches(0.18),
            cy + Inches(0.55),
            cw - Inches(0.36),
            ch - Inches(0.7),
            [lead] + bullets,
            size=12,
            color=MID,
            spacing=4,
            bold_first=True,
        )

    my = Inches(5.05)
    textbox(slide, Inches(0.35), my, Inches(8), Inches(0.28), metrics_title, size=14, bold=True, color=DARK)
    mw = Inches(2.02)
    mh = Inches(0.95)
    mg = Inches(0.12)
    for i, (label, baseline, value) in enumerate(metrics[:6]):
        x = Inches(0.35) + i * (mw + mg)
        y = my + Inches(0.32)
        rect(slide, x, y, mw, mh, fill=WHITE, line=BORDER, line_w=1)
        textbox(slide, x + Inches(0.12), y + Inches(0.06), mw - Inches(0.2), Inches(0.22), label, size=10, color=MUTED)
        textbox(slide, x + Inches(0.12), y + Inches(0.28), mw - Inches(0.2), Inches(0.26), baseline, size=10, color=MID)
        textbox(slide, x + Inches(0.12), y + Inches(0.55), mw - Inches(0.2), Inches(0.32), value, size=14, bold=True, color=ACCENT)

    fy = Inches(6.55)
    rect(slide, 0, fy, W, H - fy, fill=WHITE, line=None)
    textbox(slide, Inches(0.4), fy + Inches(0.1), Inches(12.5), Inches(0.28), action, size=12, bold=True, color=DARK)
    textbox(slide, Inches(0.4), fy + Inches(0.42), Inches(12.5), Inches(0.28), source, size=10, color=MUTED)
    return slide


def add_evidence_slide(prs, eyebrow, items):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    rect(slide, 0, 0, W, H, fill=BG)
    rect(slide, 0, 0, W, Inches(0.95), fill=WHITE, line=None)
    textbox(slide, Inches(0.4), Inches(0.16), Inches(12.5), Inches(0.28), eyebrow, size=12, color=MUTED)
    textbox(slide, Inches(0.4), Inches(0.42), Inches(12.5), Inches(0.4), "数据佐证｜官方回放截图", size=22, bold=True, color=DARK)

    n = max(len(items), 1)
    gap = Inches(0.22)
    margin = Inches(0.35)
    usable = W - 2 * margin - gap * (n - 1)
    cw = usable / n
    ch = Inches(5.55)
    cy = Inches(1.15)

    for i, (img_path, highlight, caption) in enumerate(items[:3]):
        x = margin + i * (cw + gap)
        rect(slide, x, cy, cw, ch, fill=WHITE, line=BORDER, line_w=1)
        textbox(slide, x + Inches(0.12), cy + Inches(0.1), cw - Inches(0.24), Inches(0.32), highlight, size=12, bold=True, color=ACCENT)
        ix = x + Inches(0.12)
        iy = cy + Inches(0.48)
        iw = cw - Inches(0.24)
        ih = Inches(4.2)
        path = Path(img_path)
        if path.exists():
            with Image.open(path) as im:
                aw, ah = im.size
            scale = min(float(iw) / aw, float(ih) / ah)
            dw = int(aw * scale)
            dh = int(ah * scale)
            ox = ix + (iw - dw) / 2
            oy = iy + (ih - dh) / 2
            slide.shapes.add_picture(str(path), int(ox), int(oy), width=dw, height=dh)
        else:
            textbox(slide, ix, iy, iw, Inches(0.4), f"[缺少截图: {path.name}]", size=12, color=ACCENT)
        textbox(slide, x + Inches(0.12), cy + Inches(4.85), cw - Inches(0.24), Inches(0.55), caption, size=11, color=MID)

    textbox(
        slide,
        Inches(0.4),
        Inches(6.95),
        Inches(12.5),
        Inches(0.3),
        "截图裁自云栖大会官方回放幻灯主体，作指标与结论佐证。",
        size=10,
        color=MUTED,
    )
    return slide


def clip(s: str, n: int) -> str:
    s = re.sub(r"\s+", " ", (s or "").strip())
    return s if len(s) <= n else s[: n - 1] + "…"


def bulletize(details: list[str], limit=4) -> list[str]:
    out = []
    for d in details[:limit]:
        t = d.strip()
        if not t.startswith("•") and not t.startswith("-"):
            t = "• " + t
        out.append(clip(t, 70))
    return out


def pick_metric_value(metrics: list[str]) -> str:
    if not metrics:
        return "—"
    # Prefer short numeric-looking metric
    scored = sorted(metrics, key=lambda m: (0 if re.search(r"[\d%×xX倍↑↓+\-]", m) else 1, len(m)))
    return clip(scored[0], 18)


def crop_evidence(src: Path, dst: Path) -> Path:
    dst.parent.mkdir(parents=True, exist_ok=True)
    if dst.exists() and dst.stat().st_size > 20000:
        return dst
    im = Image.open(src).convert("RGB")
    w, h = im.size
    top = int(h * 0.08)
    bot = int(h * 0.94)
    right = int(w * 0.58)
    crop = im.crop((int(w * 0.01), top, right, bot))
    if crop.width > 1400:
        nh = int(crop.height * 1400 / crop.width)
        crop = crop.resize((1400, nh), Image.Resampling.LANCZOS)
    crop.save(dst, "PNG", optimize=True)
    return dst


def load_talks(notes_dir: Path) -> list[dict]:
    talks = []
    for name in ["talks_1_5.json", "talks_6_10.json", "talks.json"]:
        path = notes_dir / name
        if not path.exists() or path.stat().st_size == 0:
            continue
        data = json.loads(path.read_text(encoding="utf-8"))
        talks.extend(data["talks"] if isinstance(data, dict) else data)
    seen = set()
    uniq = []
    for t in talks:
        if t["id"] in seen:
            continue
        seen.add(t["id"])
        uniq.append(t)
    return sorted(uniq, key=lambda x: x["id"])


def build_from_talk(talk: dict, *, forum_root: Path, source_url: str, audience_note: str = "") -> Path:
    out_dir = forum_root / "docs"
    shots = forum_root / "screenshots"
    evidence_dir = forum_root / "assets" / "evidence" / talk["id"]
    out_dir.mkdir(parents=True, exist_ok=True)
    evidence_dir.mkdir(parents=True, exist_ok=True)

    tid = talk["id"]
    title = talk.get("title", "")
    speaker = talk.get("speaker", "")
    org = talk.get("org", "")
    folder = talk["folder"]
    points = talk.get("tech_points") or []
    overview = talk.get("overview") or ""
    summaries = talk.get("summary_insights") or []

    eyebrow = f"{tid} · {speaker}" + (f" · {org}" if org else "")
    eyebrow = clip(eyebrow, 70)
    insight = "洞察：" + clip(summaries[0] if summaries else overview, 48)
    judgment = "一句话判断：" + clip(overview, 110)

    cards = []
    for i, tp in enumerate(points[:3], 1):
        details = tp.get("details") or []
        lead = clip(details[0] if details else tp.get("title", ""), 42)
        bullets = bulletize(details[1:] if len(details) > 1 else details[:1], limit=3)
        cards.append((f"{i:02d}  {clip(tp['title'], 16)}", lead, bullets))
    while len(cards) < 3:
        cards.append((f"{len(cards)+1:02d}  补充", "详见技术纪要", ["• 见 Word 展开"]))

    metrics = []
    for tp in points:
        if len(metrics) >= 6:
            break
        ms = tp.get("metrics") or []
        label = clip(tp.get("title", "指标"), 12)
        baseline = clip(ms[1], 16) if len(ms) > 1 else clip((tp.get("details") or [""])[0], 16)
        value = pick_metric_value(ms) if ms else clip((tp.get("details") or ["—"])[0], 14)
        metrics.append((label, baseline, value))
    while len(metrics) < 6:
        metrics.append(("—", "—", "—"))

    action_bits = summaries[:2] if summaries else [overview]
    insight_key = "hcs_insight" if any(tp.get("hcs_insight") for tp in points) else "competitor_insight"
    if not action_bits and points:
        action_bits = [points[0].get(insight_key) or ""]
    action = "可落地动作：" + clip("；".join(x for x in action_bits if x), 120)
    source = f"来源：{source_url} · Talk {tid} {speaker}"
    if audience_note:
        source = f"{source}  ·  视角：{audience_note}"

    # evidence: 3 cropped screenshots
    ev_items = []
    for tp in points:
        if len(ev_items) >= 3:
            break
        img = tp.get("image")
        if not img:
            continue
        src = shots / folder / img
        if not src.exists():
            continue
        dst = evidence_dir / f"{Path(img).stem}-crop.png"
        crop_evidence(src, dst)
        highlight = pick_metric_value(tp.get("metrics") or []) if tp.get("metrics") else clip(tp["title"], 20)
        if highlight == "—":
            highlight = clip(tp["title"], 20)
        ev_items.append((dst, highlight, clip(tp["title"], 28)))
    while len(ev_items) < 3 and points:
        # fallback any png in folder
        folder_pngs = sorted((shots / folder).glob("*.png"))
        for src in folder_pngs:
            if len(ev_items) >= 3:
                break
            dst = evidence_dir / f"{src.stem}-crop.png"
            crop_evidence(src, dst)
            ev_items.append((dst, "回放截图", src.name))
        break

    prs = Presentation()
    prs.slide_width = W
    prs.slide_height = H
    add_insight_slide(prs, eyebrow, insight, judgment, cards, metrics, action, source)
    add_evidence_slide(prs, eyebrow, ev_items)

    safe_title = re.sub(r'[\\/:*?"<>|]', "", title)[:28]
    out = out_dir / f"洞察一页-{tid}-{safe_title}.pptx"
    prs.save(out)
    art = Path("/opt/cursor/artifacts/yunqi-insight-ppt")
    art.mkdir(parents=True, exist_ok=True)
    prs.save(art / out.name)
    return out


FORUMS = {
    "pai": {
        "root": Path("/workspace/yunqi-2026-pai-forum"),
        "source": "https://yunqi.aliyun.com/2026/session?agendaId=123",
        "audience": "",
    },
    "agenda201": {
        "root": Path("/workspace/yunqi-2026-agenda-201"),
        "source": "https://yunqi.aliyun.com/2026/session?agendaId=201",
        "audience": "HCS 技术规划",
    },
}


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("--forum", choices=["pai", "agenda201", "all"], default="all")
    args = ap.parse_args()
    targets = list(FORUMS.keys()) if args.forum == "all" else [args.forum]
    for key in targets:
        cfg = FORUMS[key]
        talks = load_talks(cfg["root"] / "notes")
        print(f"=== {key}: {len(talks)} talks ===")
        for talk in talks:
            out = build_from_talk(
                talk,
                forum_root=cfg["root"],
                source_url=cfg["source"],
                audience_note=cfg["audience"],
            )
            print("wrote", out)


if __name__ == "__main__":
    main()
