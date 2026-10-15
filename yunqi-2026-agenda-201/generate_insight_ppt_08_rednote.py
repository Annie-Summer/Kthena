#!/usr/bin/env python3
"""ACK × 小红书 一页洞察 PPT：Cost/Token 主线 + Serving Stack + PD→RBG。"""
from __future__ import annotations

from pathlib import Path

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import PP_ALIGN
from pptx.oxml import parse_xml
from pptx.oxml.ns import qn
from pptx.util import Inches, Pt

ROOT = Path(__file__).resolve().parent
OUT = ROOT / "docs"
ART = Path("/opt/cursor/artifacts/yunqi-insight-ppt")

BG = RGBColor(0xF5, 0xF7, 0xF9)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)
DARK = RGBColor(0x1A, 0x1A, 0x1A)
MID = RGBColor(0x4A, 0x4A, 0x4A)
MUTED = RGBColor(0x6B, 0x6B, 0x6B)
BORDER = RGBColor(0xD0, 0xD8, 0xDE)
ACCENT = RGBColor(0x8B, 0x1A, 0x2B)
FONT = "Microsoft YaHei"
W = Inches(13.333)
H = Inches(7.5)


def set_run(run, size, bold=False, color=DARK):
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.color.rgb = color
    run.font.name = FONT
    rPr = run._r.get_or_add_rPr()
    ea = rPr.find(qn("a:ea"))
    if ea is None:
        ea = parse_xml(
            f'<a:ea xmlns:a="http://schemas.openxmlformats.org/drawingml/2006/main" typeface="{FONT}"/>'
        )
        rPr.append(ea)
    else:
        ea.set("typeface", FONT)


def rect(slide, x, y, w, h, fill=None, line=None, line_w=1.0):
    sh = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, x, y, w, h)
    if fill is None:
        sh.fill.background()
    else:
        sh.fill.solid()
        sh.fill.fore_color.rgb = fill
    if line is None:
        sh.line.fill.background()
    else:
        sh.line.color.rgb = line
        sh.line.width = Pt(line_w)
    return sh


def textbox(slide, x, y, w, h, text, size=14, bold=False, color=DARK):
    box = slide.shapes.add_textbox(x, y, w, h)
    tf = box.text_frame
    tf.word_wrap = True
    p = tf.paragraphs[0]
    p.alignment = PP_ALIGN.LEFT
    run = p.add_run()
    run.text = text
    set_run(run, size=size, bold=bold, color=color)
    return box


def multilines(slide, x, y, w, h, lines, size=11, color=MID, spacing=2, bold_first=False):
    box = slide.shapes.add_textbox(x, y, w, h)
    tf = box.text_frame
    tf.word_wrap = True
    for i, line in enumerate(lines):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = PP_ALIGN.LEFT
        p.space_after = Pt(spacing)
        run = p.add_run()
        run.text = line
        set_run(
            run,
            size=size,
            bold=(bold_first and i == 0),
            color=DARK if (bold_first and i == 0) else color,
        )
    return box


def main():
    prs = Presentation()
    prs.slide_width = W
    prs.slide_height = H
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    rect(slide, 0, 0, W, H, fill=BG)
    rect(slide, 0, 0, W, Inches(1.22), fill=WHITE, line=None)

    textbox(
        slide,
        Inches(0.35),
        Inches(0.08),
        Inches(12.6),
        Inches(0.24),
        "08 · 车漾 / 熊峰 · 阿里云容器服务 ACK × 小红书",
        size=11,
        color=MUTED,
    )
    textbox(
        slide,
        Inches(0.35),
        Inches(0.34),
        Inches(12.6),
        Inches(0.72),
        "从「容器服务 × 小红书」议题看 ACK 围绕 Cost/Token 的 AI 推理降本与编排演进实践",
        size=17,
        bold=True,
        color=DARK,
    )

    jy = Inches(1.32)
    rect(slide, Inches(0.3), jy, Inches(12.7), Inches(0.68), fill=WHITE, line=BORDER, line_w=1)
    textbox(
        slide,
        Inches(0.45),
        jy + Inches(0.06),
        Inches(12.4),
        Inches(0.56),
        "一句话判断：阿里云 ACK × 小红书以 Cost/Million Tokens = GPU小时成本 ÷ (Throughput × Utilization) 为北极星；"
        "用网关智能路由 + ModelScale 权重秒级加载 + InferenceKit 拓扑注入抬分母，用静态 PD→动态 RBG 演进编排，"
        "同 SLO 吞吐 2~3×、GPU 节省约 50%。",
        size=11,
        bold=True,
        color=DARK,
    )

    cards = [
        (
            "01  Serving Stack 三件套",
            "技术启示：冷启命中与拓扑正确性，比引擎峰值更能改写 Token 成本",
            [
                "• 网关：公平调度+负载/缓存亲和；TTFT -75%、ITL -40%、吞吐 3×",
                "• ModelScale：RDMA+GDR+P2P 分层权重；加载 120s→20s",
                "• InferenceKit：Webhook→Wrapper→插件；无静默、Summary 自证",
                "• NUMA 错配吞吐约 -13%~-14%；ICN64 vs RDMA 吞吐 +40%~55%",
            ],
        ),
        (
            "02  静态 PD → 动态 RBG",
            "技术启示：固定配比成组部署只够稳态，动态负载要分角色扩缩",
            [
                "• 静态成组：组内 P:D 固定，整组调度/扩缩/发布；原地升级+四级缓存",
                "• 动态分角色：Router/P/D/KV 独立扩缩与发布；按 KV 前缀命中全局选实例",
                "• 生产：同 SLO 吞吐 2~3×，GPU 节省约 50%；开源 sgl-project/rbg",
                "• 向上延伸：快照复用 + Rollout 托管，承接训推一体执行面",
            ],
        ),
    ]
    cy, ch, gap = Inches(2.14), Inches(2.68), Inches(0.22)
    n = len(cards)
    cw = (Inches(12.7) - gap * (n - 1)) / n
    for i, (title, lead, bullets) in enumerate(cards):
        x = Inches(0.3) + i * (cw + gap)
        rect(slide, x, cy, cw, ch, fill=WHITE, line=BORDER, line_w=1)
        textbox(slide, x + Inches(0.16), cy + Inches(0.1), cw - Inches(0.32), Inches(0.32), title, size=14, bold=True, color=ACCENT)
        multilines(
            slide,
            x + Inches(0.16),
            cy + Inches(0.46),
            cw - Inches(0.32),
            ch - Inches(0.56),
            [lead] + bullets,
            size=11,
            color=MID,
            spacing=2,
            bold_first=True,
        )

    metrics = [
        ("网关 TTFT", "典型场景", "-75%"),
        ("网关吞吐", "典型场景", "3×"),
        ("权重加载", "120s → 20s", "约 6×"),
        ("P2P vs OSS", "Qwen3-32B", "13.5×"),
        ("同 SLO 吞吐", "小红书生产", "2~3×"),
        ("GPU 资源", "小红书生产", "节省 50%"),
    ]
    my = Inches(4.78)
    textbox(slide, Inches(0.3), my, Inches(10), Inches(0.22), "关键指标对比与提升", size=11, bold=True, color=DARK)
    mw, mh, mg = Inches(2.05), Inches(0.78), Inches(0.1)
    for i, (label, baseline, value) in enumerate(metrics):
        x = Inches(0.3) + i * (mw + mg)
        y = my + Inches(0.22)
        rect(slide, x, y, mw, mh, fill=WHITE, line=BORDER, line_w=1)
        textbox(slide, x + Inches(0.08), y + Inches(0.02), mw - Inches(0.14), Inches(0.18), label, size=9, color=MUTED)
        textbox(slide, x + Inches(0.08), y + Inches(0.2), mw - Inches(0.14), Inches(0.2), baseline, size=9, color=MID)
        textbox(slide, x + Inches(0.08), y + Inches(0.4), mw - Inches(0.14), Inches(0.32), value, size=12, bold=True, color=ACCENT)

    fy = Inches(5.9)
    rect(slide, 0, fy, W, H - fy, fill=WHITE, line=None)
    multilines(
        slide,
        Inches(0.35),
        fy + Inches(0.04),
        Inches(12.6),
        Inches(1.2),
        [
            "技术启示：",
            "• 发布准入与弹性观测指标集：TTFT、吞吐、权重分层命中率、冷启就绪时延、NUMA 与卡间互联拓扑正确性，形成可闭环验收。",
            "• ModelScale 冷温热权重分层（先 OSS，再 Mooncake，再到 FlashGPU/P2P），配合 InferenceKit 拓扑注入（NUMA 与互联），避免仅做引擎并行或 batching 时收益被冷启与错配抵消。",
            "• 编排上，固定 P:D 成组部署只适合稳态生产；要支撑动态负载，需按 Router、Prefill、Decode、KV 分角色独立扩缩、协同发布，并按 KV 前缀命中在全局实例池中调度请求；Agent RL 的环境快照断点续跑与轨迹托管也跑在同一套角色编排与弹性控制面上。",
        ],
        size=11,
        color=DARK,
        spacing=2,
        bold_first=True,
    )
    textbox(
        slide,
        Inches(0.35),
        fy + Inches(1.24),
        Inches(12.6),
        Inches(0.22),
        "来源：云栖大会 agendaId=201｜容器服务 × 小红书：大模型推理与 Agent RL · 车漾 / 熊峰",
        size=9,
        color=MUTED,
    )

    OUT.mkdir(parents=True, exist_ok=True)
    ART.mkdir(parents=True, exist_ok=True)
    out = OUT / "洞察一页-08-容器服务 × 小红书：大模型推理与 Agent RL.pptx"
    prs.save(out)
    prs.save(ART / out.name)
    print("wrote", out, "slides=1")


if __name__ == "__main__":
    main()
