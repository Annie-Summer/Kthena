#!/usr/bin/env python3
"""Generate Word docs for Yunqi 2026 agendaId=201 Agent Infra forum.

Audience lens: HCS 技术规划团队 — competitor/peer implications framed for
Huawei Cloud Stack technical planning (architecture, product roadmap, gaps).
"""
from __future__ import annotations

import json
from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor

ROOT = Path(__file__).resolve().parent
NOTES = ROOT / "notes"
DOCS = ROOT / "docs"
SHOTS = ROOT / "screenshots"
SOURCE = "https://yunqi.aliyun.com/2026/session?agendaId=201"
AUDIENCE = "HCS 技术规划团队"


def set_run_font(run, name="微软雅黑", size=11, bold=False, color=None):
    run.font.name = name
    run._element.rPr.rFonts.set(qn("w:eastAsia"), name)
    run.font.size = Pt(size)
    run.bold = bold
    if color:
        run.font.color.rgb = RGBColor(*color)


def add_heading(doc, text, level=1):
    p = doc.add_heading(text, level=level)
    for run in p.runs:
        set_run_font(run, size=16 if level == 1 else 13, bold=True)
    return p


def add_para(doc, text, *, bold=False, size=11, space_after=6):
    p = doc.add_paragraph()
    run = p.add_run(text)
    set_run_font(run, size=size, bold=bold)
    p.paragraph_format.space_after = Pt(space_after)
    p.paragraph_format.line_spacing = 1.25
    return p


def add_bullets(doc, items):
    for item in items:
        p = doc.add_paragraph(style="List Bullet")
        run = p.add_run(item)
        set_run_font(run, size=11)
        p.paragraph_format.space_after = Pt(3)


def add_image(doc, path: Path, caption: str | None = None, width=5.6):
    if not path.exists():
        add_para(doc, f"[缺少截图: {path.name}]", bold=True)
        return
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run()
    run.add_picture(str(path), width=Inches(width))
    if caption:
        cp = doc.add_paragraph()
        cp.alignment = WD_ALIGN_PARAGRAPH.CENTER
        r = cp.add_run(caption)
        set_run_font(r, size=9, color=(90, 90, 90))


def _slug(s: str) -> str:
    bad = '/\\:*?"<>|'
    for c in bad:
        s = s.replace(c, "")
    return s.replace(" ", "").replace("：", "-").replace(":", "-")[:40]


def build_talk_doc(talk: dict) -> Path:
    doc = Document()
    section = doc.sections[0]
    section.top_margin = Inches(0.7)
    section.bottom_margin = Inches(0.7)
    section.left_margin = Inches(0.8)
    section.right_margin = Inches(0.8)

    title = talk["title"]
    speaker = talk.get("speaker", "")
    org = talk.get("org", "")
    folder = talk["folder"]
    tid = talk["id"]

    add_heading(doc, f"{tid} · {title}", 1)
    add_para(doc, f"演讲：{speaker}" + (f" · {org}" if org else ""), size=11, bold=True)
    add_para(doc, f"阅读视角：{AUDIENCE}", size=10, bold=True)
    add_para(doc, f"来源回放：{SOURCE}", size=9, space_after=10)

    add_heading(doc, "一、议题概览", 2)
    add_para(doc, talk.get("overview", ""))

    add_heading(doc, "二、关键技术细节（附图佐证）", 2)
    add_para(
        doc,
        "以下按技术点展开：技术细节与指标 → 官方回放截图佐证 → 对 HCS 技术规划的启示（架构/产品/差距/可落项）。",
        size=10,
    )

    for i, tp in enumerate(talk.get("tech_points", []), 1):
        add_heading(doc, f"{i}. {tp['title']}", 3)
        details = tp.get("details") or []
        if details:
            add_para(doc, "技术细节", bold=True, size=11, space_after=3)
            add_bullets(doc, details)
        metrics = tp.get("metrics") or []
        if metrics:
            add_para(doc, "关键指标", bold=True, size=11, space_after=3)
            add_bullets(doc, metrics)
        img_name = tp.get("image")
        if img_name:
            add_image(doc, SHOTS / folder / img_name, caption=f"图示佐证：{tp['title']}（{img_name}）")
        insight = tp.get("hcs_insight") or tp.get("competitor_insight")
        if insight:
            add_para(doc, "对 HCS 技术规划的启示", bold=True, size=11, space_after=3)
            add_para(doc, insight, size=11, space_after=10)

    add_heading(doc, "三、综合启示（HCS 技术规划）", 2)
    insights = talk.get("summary_insights") or []
    if insights:
        add_bullets(doc, insights)
    else:
        add_para(doc, "（见各技术点附带启示）")

    add_para(doc, "", space_after=4)
    add_para(
        doc,
        "说明：截图来自云栖大会官方回放；指标以演讲幻灯口径为准。启示面向 HCS 技术规划的对标与能力补齐，不代表官方立场。",
        size=9,
        space_after=2,
    )

    DOCS.mkdir(parents=True, exist_ok=True)
    safe_speaker = (speaker.split("，")[0].split("/")[0].split("、")[0]).strip() or "speaker"
    out = DOCS / f"{tid}-{_slug(title)}-{_slug(safe_speaker)}.docx"
    doc.save(out)
    return out


def load_talks() -> list[dict]:
    talks = []
    for name in ["talks.json", "talks_1_5.json", "talks_6_10.json"]:
        path = NOTES / name
        if not path.exists() or path.stat().st_size == 0:
            continue
        data = json.loads(path.read_text(encoding="utf-8"))
        if isinstance(data, dict) and "talks" in data:
            talks.extend(data["talks"])
        elif isinstance(data, list):
            talks.extend(data)
    seen = set()
    uniq = []
    for t in talks:
        if t["id"] in seen:
            continue
        seen.add(t["id"])
        uniq.append(t)
    return sorted(uniq, key=lambda x: x["id"])


def main():
    talks = load_talks()
    if not talks:
        raise SystemExit("No talk notes found under notes/")
    outs = []
    for talk in talks:
        out = build_talk_doc(talk)
        outs.append(out)
        print("wrote", out)
    print(f"total {len(outs)} docs")


if __name__ == "__main__":
    main()
