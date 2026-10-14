#!/usr/bin/env python3
"""Agentic Inference 洞察 PPT：按.md全量技术点覆盖（2页洞察）。"""
from __future__ import annotations

from pathlib import Path

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import PP_ALIGN
from pptx.oxml import parse_xml
from pptx.oxml.ns import qn
from pptx.util import Inches, Pt

ROOT = Path(__file__).resolve().parent
OUT = ROOT / "docs"
ART = Path("/opt/cursor/artifacts/yunqi-insight-ppt")

BG = RGBColor(0xF5, 0xF7, 0xF9)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)
DARK = RGBColor(0x1A, 0x1A, 0x1A)
MID = RGBColor(0x4A, 0x4A, 0x4A)
MUTED = RGBColor(0x6B, 0x6B, 0x6B)
BORDER = RGBColor(0xD0, 0xD8, 0xDE)
ACCENT = RGBColor(0x8B, 0x1A, 0x2B)
FONT = "Microsoft YaHei"
W = Inches(13.333)
H = Inches(7.5)


def set_run(run, size, bold=False, color=DARK):
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.color.rgb = color
    run.font.name = FONT
    rPr = run._r.get_or_add_rPr()
    ea = rPr.find(qn("a:ea"))
    if ea is None:
        ea = parse_xml(
            f'<a:ea xmlns:a="http://schemas.openxmlformats.org/drawingml/2006/main" typeface="{FONT}"/>'
        )
        rPr.append(ea)
    else:
        ea.set("typeface", FONT)


def rect(slide, x, y, w, h, fill=None, line=None, line_w=1.0):
    sh = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, x, y, w, h)
    if fill is None:
        sh.fill.background()
    else:
        sh.fill.solid()
        sh.fill.fore_color.rgb = fill
    if line is None:
        sh.line.fill.background()
    else:
        sh.line.color.rgb = line
        sh.line.width = Pt(line_w)
    return sh


def textbox(slide, x, y, w, h, text, size=14, bold=False, color=DARK):
    box = slide.shapes.add_textbox(x, y, w, h)
    tf = box.text_frame
    tf.word_wrap = True
    p = tf.paragraphs[0]
    p.alignment = PP_ALIGN.LEFT
    run = p.add_run()
    run.text = text
    set_run(run, size=size, bold=bold, color=color)
    return box


def multilines(slide, x, y, w, h, lines, size=11, color=MID, spacing=2, bold_first=False):
    box = slide.shapes.add_textbox(x, y, w, h)
    tf = box.text_frame
    tf.word_wrap = True
    for i, line in enumerate(lines):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = PP_ALIGN.LEFT
        p.space_after = Pt(spacing)
        run = p.add_run()
        run.text = line
        set_run(
            run,
            size=size,
            bold=(bold_first and i == 0),
            color=DARK if (bold_first and i == 0) else color,
        )
    return box


def add_cards(slide, cards, cy, ch=Inches(2.55)):
    cw, gap = Inches(4.1), Inches(0.18)
    for i, (title, lead, bullets) in enumerate(cards):
        x = Inches(0.3) + i * (cw + gap)
        rect(slide, x, cy, cw, ch, fill=WHITE, line=BORDER, line_w=1)
        textbox(slide, x + Inches(0.12), cy + Inches(0.08), cw - Inches(0.24), Inches(0.3), title, size=13, bold=True, color=ACCENT)
        multilines(
            slide,
            x + Inches(0.12),
            cy + Inches(0.4),
            cw - Inches(0.24),
            ch - Inches(0.5),
            [lead] + bullets,
            size=10,
            color=MID,
            spacing=1,
            bold_first=True,
        )


def add_metrics(slide, metrics, my, title="关键指标对比与提升"):
    textbox(slide, Inches(0.3), my, Inches(10), Inches(0.24), title, size=12, bold=True, color=DARK)
    mw, mh, mg = Inches(2.05), Inches(0.95), Inches(0.1)
    for i, (label, baseline, value) in enumerate(metrics[:6]):
        x = Inches(0.3) + i * (mw + mg)
        y = my + Inches(0.26)
        rect(slide, x, y, mw, mh, fill=WHITE, line=BORDER, line_w=1)
        textbox(slide, x + Inches(0.08), y + Inches(0.04), mw - Inches(0.14), Inches(0.2), label, size=9, color=MUTED)
        textbox(slide, x + Inches(0.08), y + Inches(0.24), mw - Inches(0.14), Inches(0.26), baseline, size=9, color=MID)
        textbox(slide, x + Inches(0.08), y + Inches(0.5), mw - Inches(0.14), Inches(0.38), value, size=12, bold=True, color=ACCENT)


def add_page1(prs):
    """覆盖：Agentic定义/四特征、三层架构、五层路由、调度策略、智能选模。"""
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    rect(slide, 0, 0, W, H, fill=BG)
    rect(slide, 0, 0, W, Inches(1.2), fill=WHITE, line=None)

    textbox(slide, Inches(0.35), Inches(0.08), Inches(12.6), Inches(0.24),
            "05 · 李文鹏 · 阿里云 PAI｜第 1/2 页", size=11, color=MUTED)
    textbox(slide, Inches(0.35), Inches(0.34), Inches(12.6), Inches(0.72),
            "从「Agentic Inference：推理系统的全链路优化」看阿里云 PAI 围绕有状态推理全链路优化的创新实践",
            size=16, bold=True, color=DARK)

    jy = Inches(1.28)
    rect(slide, Inches(0.3), jy, Inches(12.7), Inches(0.62), fill=WHITE, line=BORDER, line_w=1)
    textbox(slide, Inches(0.45), jy + Inches(0.06), Inches(12.4), Inches(0.5),
            "一句话判断：Agent 多轮长上下文下，无状态 QPS 调度保不住前缀命中与 Token 成本；"
            "PAI 以 Session/Cache-Aware 调度叠加难度分流选模，把 KV 命中率做到约 79%、TTFT 降约 33%、TPS 约翻倍，"
            "并用 Auto 选模在质量提升同时把成本指数压到约 35%。",
            size=11, bold=True, color=DARK)

    cards = [
        (
            "01  Agentic 负载四特征",
            "技术启示：调度 SLO 要从实例 QPS，切到 Session 命中与 Token 成本",
            [
                "• 长上下文：百→数万 Token；4K 请求 KV 可超 200MB",
                "• KV 高命中：多轮前缀重叠，命中率决定 TTFT/吞吐",
                "• Token 级规划：长请求成本远高于短请求，禁只用 QPS",
                "• 有状态调度：同会话回原实例；伸缩须保持 KV 连续",
            ],
        ),
        (
            "02  三层架构 × 五层路由",
            "技术启示：调度、执行、全局 KV 必须同代交付，单点引擎优化不够",
            [
                "• 调度层：统一 API、热更新治理、容错、智能择址、实例角色",
                "• 执行层：全形态 PD、TP/EP/DP 感知、vLLM/SGLang/自定义",
                "• 缓存层：DRAM+SSD 池、RDMA 直通、L1/L2/L3 命中可视",
                "• 路由五层：入口→治理→容错→调度→Normal/P/D 实例",
            ],
        ),
        (
            "03  亲和调度 + 智能选模",
            "技术启示：先保会话/前缀命中，再用难度分流固化质量-成本边界",
            [
                "• Session-Aware(300)：strict/soft，过载熔断回退",
                "• Cache-Aware(200)：块哈希索引；Least-Token(0) 兜底",
                "• PD/DP Rank 差异化：P 追前缀，D 追负载，MoE 防局部过热",
                "• Auto 选模 L1/L2/L3：客户端写 auto，网关实例自动切候选",
            ],
        ),
    ]
    add_cards(slide, cards, Inches(2.02), ch=Inches(2.72))

    add_metrics(
        slide,
        [
            ("Cache 命中率", "基线→Cache-Aware", "38.65%→78.62%"),
            ("TTFT", "Cache-Aware", "-33.3%"),
            ("TPS", "Cache-Aware", "+100%"),
            ("选模质量", "纯L1→Auto", "47.3%→82.7%"),
            ("选模成本指数", "纯L3→Auto", "100→35.1(-65%)"),
            ("流量承接", "L1+L2 占比", "85%+"),
        ],
        Inches(4.78),
    )

    fy = Inches(6.18)
    rect(slide, 0, fy, W, H - fy, fill=WHITE, line=None)
    textbox(slide, Inches(0.35), fy + Inches(0.05), Inches(12.6), Inches(0.72),
            "技术启示：Agent 多轮负载下，发布准入与弹性观测不能只看实例 QPS/利用率，"
            "还需把 Session 亲和成功率、KV 前缀命中率、TPM/Token 配额一并纳入；"
            "否则只在推理引擎（并行策略、batching）上拿到的 TTFT/TPS 收益，会被请求打散后的前缀未命中吞掉。"
            "控制面要以 Session→Cache→Least-Token 策略级联择址，并与 PD/并行感知执行面、全局 KV 同代交付；"
            "选模侧用难度分流，让旗舰模型只承接难题流量。",
            size=11, bold=True, color=DARK)
    textbox(slide, Inches(0.35), fy + Inches(0.78), Inches(12.6), Inches(0.24),
            "来源：云栖大会｜阿里云 PAI《Agentic Inference：推理系统的全链路优化》· 覆盖第1–5章",
            size=9, color=MUTED)


def add_page2(prs):
    """覆盖：联邦、可观测、流控、灰度、批量、录制回放、引擎PD、FlashKV、模型分发、机密计算。"""
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    rect(slide, 0, 0, W, H, fill=BG)
    rect(slide, 0, 0, W, Inches(1.12), fill=WHITE, line=None)

    textbox(slide, Inches(0.35), Inches(0.08), Inches(12.6), Inches(0.24),
            "05 · 续 · 平台能力闭环｜第 2/2 页", size=11, color=MUTED)
    textbox(slide, Inches(0.35), Inches(0.34), Inches(12.6), Inches(0.64),
            "从「联邦·可观测·治理·引擎·FlashKV·分发·机密计算」看 PAI 将 Agentic 推理做成可运营平台能力",
            size=16, bold=True, color=DARK)

    jy = Inches(1.18)
    rect(slide, Inches(0.3), jy, Inches(12.7), Inches(0.58), fill=WHITE, line=BORDER, line_w=1)
    textbox(slide, Inches(0.45), jy + Inches(0.06), Inches(12.4), Inches(0.48),
            "一句话判断：有状态调度之上，PAI 把跨 Region 库存联邦、Token 计量闸门、可用性感知灰度与录制回放做成控制面默认能力，"
            "再用 FlashKV/CacheFS 把命中率推到约 94%、模型加载压到秒级，并以 TEE 远程证明把安全做成可验证执行。",
            size=11, bold=True, color=DARK)

    cards = [
        (
            "04  联邦·观测·流控·灰度",
            "技术启示：库存、计量、发布必须进入控制面，不能只做引擎侧优化",
            [
                "• 联邦：PRIMARY 全局 Watch；成员区按实例 ID 透明转发",
                "• 观测：TTFT/TPOT/排队分位 + L1/L2/L3 命中 + 实例画像",
                "• 流控：Coding/Token Plan；时段/并发/RPM·TPM/周期四闸门",
                "• 灰度：权重即生命周期；可用性感知；PD 配对完整性校验",
            ],
        ),
        (
            "05  批量·回放·引擎 PD",
            "技术启示：真实流量验收 + 全形态 PD 统一接入，才能压住发布风险",
            [
                "• 批量：≤5万行/200MB；高峰让路、夜间吃尽闲置 GPU",
                "• 录制回放：字节级保真 + 原节奏；发布前用真实流量验收",
                "• PD 三模式：Prefill代理 / 网关调度 / Decode主导 全兼容",
                "• 并行感知：P→TP；D→DP Rank；MoE 防局部过热",
            ],
        ),
        (
            "06  FlashKV·分发·机密计算",
            "技术启示：缩容后命中与冷启时延，决定弹性是否真能省 GPU",
            [
                "• FlashKV：引擎只留活跃 KV；DRAM/SSD + RDMA 零拷贝",
                "• 命中再升：78.62%→94.32%；TTFT -25%；TPS +40%",
                "• CacheFS/CacheRoot：P2P+中心缓存，加载数十分钟→秒级",
                "• PAI-CC：本地加密上云→远程证明放行→TEE 内端到端密态",
            ],
        ),
    ]
    add_cards(slide, cards, Inches(1.9), ch=Inches(2.78))

    add_metrics(
        slide,
        [
            ("FlashKV 命中率", "叠加全局 KV", "78.62%→94.32%"),
            ("FlashKV TTFT", "相对前一阶段", "-25%"),
            ("FlashKV TPS", "相对前一阶段", "+40%"),
            ("模型加载", "远端直读→分发加速", "数十分钟→秒级"),
            ("批量规格", "单任务上限", "5万行 / 200MB"),
            ("安全形态", "承诺式→可验证", "TEE+远程证明"),
        ],
        Inches(4.82),
        title="平台能力关键指标",
    )

    fy = Inches(6.28)
    rect(slide, 0, fy, W, H - fy, fill=WHITE, line=None)
    textbox(slide, Inches(0.35), fy + Inches(0.06), Inches(12.6), Inches(0.58),
            "技术启示：跨 Region 联邦与 Token/Coding 四维闸门要把库存利用率与计量写成准入默认项；"
            "灰度需可用性感知并校验 PD 配对完整性，发布前用录制回放做真实流量验收。"
            "弹性侧要把活跃 KV 外置到 DRAM/SSD（FlashKV）并把权重分发做到秒级就绪，"
            "否则缩容后再扩会重新支付冷启与未命中成本；安全侧以 TEE+远程证明替代口头合规承诺。",
            size=11, bold=True, color=DARK)
    textbox(slide, Inches(0.35), fy + Inches(0.6), Inches(12.6), Inches(0.26),
            "来源：云栖大会｜阿里云 PAI《Agentic Inference：推理系统的全链路优化》· 覆盖第6–15章",
            size=9, color=MUTED)


def main():
    prs = Presentation()
    prs.slide_width = W
    prs.slide_height = H
    add_page1(prs)
    add_page2(prs)

    OUT.mkdir(parents=True, exist_ok=True)
    ART.mkdir(parents=True, exist_ok=True)
    out = OUT / "洞察一页-05-Agentic Inference：推理系统的全链路优化.pptx"
    out2 = OUT / "洞察一页-05-Agentic Inference：从议题看有状态推理全链路优化.pptx"
    prs.save(out)
    prs.save(out2)
    prs.save(ART / out.name)
    print("wrote", out, "slides=", len(prs.slides))


if __name__ == "__main__":
    main()
